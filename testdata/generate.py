"""
Генератор синтетических тестовых документов, expected.json и docs/RULES.md.

Все имена, реквизиты и организации вымышлены. ИНН начинается с кода региона 00,
которого не существует, поэтому совпасть с настоящей организацией не может.

Запуск из корня репозитория:
    pip install python-docx==1.2.0 reportlab==5.0.1
    python testdata/generate.py

Скрипт перед записью файлов проверяет ожидаемые замечания той же логикой,
что у DefaultRuleEngine (пустое значение проверяет только REQUIRED, дата —
строгий разбор по шаблону, шаблон — совпадение всей строки). Расхождение —
ошибка генерации, а не документа.
"""
import json
import re
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent
TESTDATA = ROOT / "testdata"
DOCUMENTS = TESTDATA / "documents"
RULES = json.loads((TESTDATA / "rules.json").read_text(encoding="utf-8"))

# Дата, относительно которой считаются ожидаемые замечания DATE_NOT_FUTURE.
# Тесты движка передают её в ValidationContext, а не берут часы машины.
REFERENCE_DATE = date(2026, 9, 21)

FOOTER = "Синтетический тестовый документ. Все имена, организации и реквизиты вымышлены."
FIXED_TIME = datetime(2026, 9, 21, 12, 0, 0)

# Шрифт с кириллицей, свободная лицензия. Ищется в системе; в репозиторий не кладётся.
FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/Library/Fonts/DejaVuSans.ttf"),
]

# ---------------------------------------------------------------------------
# Документы. lines — строки документа по порядку; fields — что должно получиться
# после идеального извлечения; issues — какие правила обязаны сработать.
# ---------------------------------------------------------------------------

MEMO_BODY = (
    "В связи с расширением отдела продаж с 1 октября 2026 года прошу согласовать "
    "закупку пяти ноутбуков для новых сотрудников. Ориентировочная стоимость — "
    "450 000 рублей, средства предусмотрены бюджетом отдела на IV квартал."
)

DOCS = [
    {
        "file": "01_memo_ok.docx",
        "type": "OFFICIAL_MEMO",
        "about": "Служебная записка без замечаний.",
        "lines": [
            ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
            ("", ""),
            ("bold", "СЛУЖЕБНАЯ ЗАПИСКА"),
            ("", "15.09.2026 № СЗ-117"),
            ("", ""),
            ("", "О закупке ноутбуков для отдела продаж"),
            ("", ""),
            ("", MEMO_BODY),
            ("", ""),
            ("", "Руководитель отдела продаж"),
            ("", "А.С. Петрова"),
        ],
        "fields": {
            "addressee": "Генеральному директору ООО «Демо-компания» Орлову В.П.",
            "author_name": "А.С. Петрова",
            "doc_date": "15.09.2026",
            "reg_number": "СЗ-117",
            "subject": "О закупке ноутбуков для отдела продаж",
            "body": MEMO_BODY,
            "signer_position": "Руководитель отдела продаж",
            "signer_name": "А.С. Петрова",
        },
        "issues": [],
    },
    {
        "file": "02_memo_errors.docx",
        "type": "OFFICIAL_MEMO",
        "about": "Нет адресата, дата в будущем, заголовок без «О», в подписи полное имя.",
        "lines": [
            ("bold", "СЛУЖЕБНАЯ ЗАПИСКА"),
            ("", "31.12.2099 № СЗ-118"),
            ("", ""),
            ("", "Закупка ноутбуков для отдела продаж"),
            ("", ""),
            ("", MEMO_BODY),
            ("", ""),
            ("", "Руководитель отдела продаж"),
            ("", "Анна Петрова"),
        ],
        "fields": {
            "author_name": "Анна Петрова",
            "doc_date": "31.12.2099",
            "reg_number": "СЗ-118",
            "subject": "Закупка ноутбуков для отдела продаж",
            "body": MEMO_BODY,
            "signer_position": "Руководитель отдела продаж",
            "signer_name": "Анна Петрова",
        },
        "issues": ["M1", "M4", "M6", "M10"],
    },
    {
        "file": "03_memo_word_date.pdf",
        "type": "OFFICIAL_MEMO",
        "about": (
            "PDF. Дата словами, номер без индекса, короткий текст, в подписи нет должности. "
            "Дата словами даёт и M3, и M4: движок не разбирает её по шаблону ДД.ММ.ГГГГ."
        ),
        "lines": [
            ("right", "Главному бухгалтеру ООО «Демо-компания» Соколовой Е.Н."),
            ("", ""),
            ("bold", "СЛУЖЕБНАЯ ЗАПИСКА"),
            ("", "15 сентября 2026 г. № 119"),
            ("", ""),
            ("", "О возмещении командировочных расходов"),
            ("", ""),
            ("", "Прошу возместить расходы."),
            ("", ""),
            ("", "Д.В. Смирнов"),
        ],
        "fields": {
            "addressee": "Главному бухгалтеру ООО «Демо-компания» Соколовой Е.Н.",
            "author_name": "Д.В. Смирнов",
            "doc_date": "15 сентября 2026 г.",
            "reg_number": "119",
            "subject": "О возмещении командировочных расходов",
            "body": "Прошу возместить расходы.",
            "signer_name": "Д.В. Смирнов",
        },
        "issues": ["M3", "M4", "M5", "M8", "M9"],
    },
    {
        "file": "04_vacation_ok.docx",
        "type": "VACATION_REQUEST",
        "about": "Заявление на отпуск без замечаний.",
        "lines": [
            ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
            ("right", "от менеджера по продажам Кузнецова Игоря Олеговича"),
            ("", ""),
            ("bold", "ЗАЯВЛЕНИЕ"),
            ("", ""),
            ("", "Прошу предоставить мне ежегодный оплачиваемый отпуск с 01.10.2026 "
                 "продолжительностью 14 календарных дней."),
            ("", ""),
            ("", "15.09.2026        И.О. Кузнецов"),
        ],
        "fields": {
            "employee_name": "Кузнецов Игорь Олегович",
            "employee_position": "менеджер по продажам",
            "vacation_type": "ежегодный оплачиваемый",
            "start_date": "01.10.2026",
            "days_count": "14",
        },
        "issues": [],
    },
    {
        "file": "05_vacation_errors.pdf",
        "type": "VACATION_REQUEST",
        "about": "PDF. Учебный отпуск, дата начала без года, дни словами.",
        "lines": [
            ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
            ("right", "от менеджера по продажам Кузнецова Игоря Олеговича"),
            ("", ""),
            ("bold", "ЗАЯВЛЕНИЕ"),
            ("", ""),
            ("", "Прошу предоставить мне учебный отпуск с 1 октября "
                 "продолжительностью двадцать восемь календарных дней."),
            ("", ""),
            ("", "15.09.2026        И.О. Кузнецов"),
        ],
        "fields": {
            "employee_name": "Кузнецов Игорь Олегович",
            "employee_position": "менеджер по продажам",
            "vacation_type": "учебный",
            "start_date": "1 октября",
            "days_count": "двадцать восемь",
        },
        "issues": ["V2", "V4", "V5"],
    },
    {
        "file": "06_support_ok.pdf",
        "type": "SUPPORT_MEASURE_REQUEST",
        "about": "PDF. Заявка на меру поддержки без замечаний.",
        "lines": [
            ("right", "В центр поддержки предпринимательства (вымышленный)"),
            ("", ""),
            ("bold", "ЗАЯВКА НА ПОЛУЧЕНИЕ МЕРЫ ПОДДЕРЖКИ"),
            ("", ""),
            ("", "Заявитель: ООО «Демо-компания»"),
            ("", "ИНН: 0012345678"),
            ("", "Категория субъекта МСП: малое предприятие"),
            ("", "Вид поддержки: консультационная поддержка по вопросам бухгалтерского учёта"),
            ("", ""),
            ("", "Заявитель подтверждает, что не является кредитной организацией, страховой "
                 "организацией, инвестиционным фондом, негосударственным пенсионным фондом, "
                 "профессиональным участником рынка ценных бумаг, ломбардом, участником "
                 "соглашения о разделе продукции, не осуществляет деятельность в сфере "
                 "игорного бизнеса и не является нерезидентом Российской Федерации."),
            ("", "Заявитель подтверждает, что аналогичная поддержка ему ранее не оказывалась."),
            ("", ""),
            ("", "Приложения: 1. Описание запроса на консультацию; 2. Копия доверенности представителя."),
            ("", ""),
            ("", "Генеральный директор        В.П. Орлов        15.09.2026"),
        ],
        "fields": {
            "applicant_name": "ООО «Демо-компания»",
            "applicant_inn": "0012345678",
            "smb_category": "малое предприятие",
            "support_type": "консультационная поддержка по вопросам бухгалтерского учёта",
            "exclusions_declaration": (
                "Заявитель подтверждает, что не является кредитной организацией, страховой "
                "организацией, инвестиционным фондом, негосударственным пенсионным фондом, "
                "профессиональным участником рынка ценных бумаг, ломбардом, участником "
                "соглашения о разделе продукции, не осуществляет деятельность в сфере "
                "игорного бизнеса и не является нерезидентом Российской Федерации."
            ),
            "no_similar_support_declaration": "Заявитель подтверждает, что аналогичная поддержка ему ранее не оказывалась.",
            "attachments_list": "Описание запроса на консультацию; Копия доверенности представителя",
        },
        "issues": [],
    },
    {
        "file": "07_support_errors.docx",
        "type": "SUPPORT_MEASURE_REQUEST",
        "about": "ИНН из 9 цифр, форма вместо категории, нет заявления об аналогичной поддержке и перечня приложений.",
        "lines": [
            ("right", "В центр поддержки предпринимательства (вымышленный)"),
            ("", ""),
            ("bold", "ЗАЯВКА НА ПОЛУЧЕНИЕ МЕРЫ ПОДДЕРЖКИ"),
            ("", ""),
            ("", "Заявитель: ООО «Демо-компания»"),
            ("", "ИНН: 001234567"),
            ("", "Категория субъекта МСП: индивидуальный предприниматель"),
            ("", "Вид поддержки: консультационная поддержка по вопросам бухгалтерского учёта"),
            ("", ""),
            ("", "Заявитель подтверждает, что не является кредитной организацией, страховой "
                 "организацией, инвестиционным фондом, негосударственным пенсионным фондом, "
                 "профессиональным участником рынка ценных бумаг, ломбардом, участником "
                 "соглашения о разделе продукции, не осуществляет деятельность в сфере "
                 "игорного бизнеса и не является нерезидентом Российской Федерации."),
            ("", ""),
            ("", "Генеральный директор        В.П. Орлов        15.09.2026"),
        ],
        "fields": {
            "applicant_name": "ООО «Демо-компания»",
            "applicant_inn": "001234567",
            "smb_category": "индивидуальный предприниматель",
            "support_type": "консультационная поддержка по вопросам бухгалтерского учёта",
            "exclusions_declaration": (
                "Заявитель подтверждает, что не является кредитной организацией, страховой "
                "организацией, инвестиционным фондом, негосударственным пенсионным фондом, "
                "профессиональным участником рынка ценных бумаг, ломбардом, участником "
                "соглашения о разделе продукции, не осуществляет деятельность в сфере "
                "игорного бизнеса и не является нерезидентом Российской Федерации."
            ),
        },
        "issues": ["S2", "S3", "S5", "S6"],
    },
]


# ---------------------------------------------------------------------------
# Самопроверка: та же семантика, что у DefaultRuleEngine
# ---------------------------------------------------------------------------

def java_date_pattern_to_regex(pattern: str) -> str:
    if pattern != "dd.MM.uuuu":
        raise ValueError(f"Самопроверка знает только dd.MM.uuuu, а не {pattern}")
    return r"\d{2}\.\d{2}\.\d{4}"


def parse_date(value: str, pattern: str):
    if not re.fullmatch(java_date_pattern_to_regex(pattern), value):
        return None
    day, month, year = (int(part) for part in value.split("."))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def violates(rule: dict, value) -> bool:
    missing = value is None or value.strip() == ""
    check = rule["check"]
    if check == "REQUIRED":
        return missing
    if missing:
        return False
    if check == "DATE_FORMAT":
        return parse_date(value, rule["expected"]) is None
    if check == "DATE_NOT_FUTURE":
        parsed = parse_date(value, rule["expected"])
        return parsed is None or parsed > REFERENCE_DATE
    if check == "MATCHES_PATTERN":
        return re.fullmatch(rule["expected"], value) is None
    if check == "MIN_LENGTH":
        return len(value) < int(rule["expected"])
    if check == "ONE_OF":
        return value not in json.loads(rule["expected"])
    raise ValueError(f"Неизвестная проверка {check}")


def self_check() -> None:
    types = {t["code"]: t for t in RULES["types"]}
    for doc in DOCS:
        doc_type = types[doc["type"]]
        field_names = {f["name"] for f in doc_type["fields"]}
        unknown = set(doc["fields"]) - field_names
        assert not unknown, f"{doc['file']}: неизвестные поля {unknown}"
        actual = [r["code"] for r in doc_type["rules"] if violates(r, doc["fields"].get(r["field"]))]
        assert actual == doc["issues"], f"{doc['file']}: ожидали {doc['issues']}, движок даст {actual}"
        joined = "\n".join(text for _, text in doc["lines"])
        for name, value in doc["fields"].items():
            probe = value.split("; ")[0] if name == "attachments_list" else value
            assert probe in joined or name in {"employee_name", "employee_position"}, (
                f"{doc['file']}: значения поля {name} нет в тексте документа"
            )


# ---------------------------------------------------------------------------
# Файлы
# ---------------------------------------------------------------------------

def write_docx(doc: dict, path: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for kind, text in doc["lines"]:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        if kind == "bold":
            run.bold = True
            paragraph.alignment = 1
        elif kind == "right":
            paragraph.alignment = 2
    document.add_paragraph()
    footer = document.add_paragraph().add_run(FOOTER)
    footer.italic = True
    footer.font.size = Pt(9)
    props = document.core_properties
    props.author = "Сибирский Вайб — тестовые данные"
    props.title = doc["about"]
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.last_modified_by = "generate.py"
    props.revision = 1
    document.save(path)


def register_font() -> str:
    for candidate in FONT_CANDIDATES:
        if candidate.exists():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(candidate)))
            return "DejaVuSans"
    raise FileNotFoundError("Нужен шрифт DejaVuSans.ttf с кириллицей")


def write_pdf(doc: dict, path: Path, font: str) -> None:
    base = ParagraphStyle("base", fontName=font, fontSize=11, leading=15)
    styles = {
        "": base,
        "bold": ParagraphStyle("bold", parent=base, alignment=1, fontSize=12),
        "right": ParagraphStyle("right", parent=base, alignment=2),
    }
    small = ParagraphStyle("small", parent=base, fontSize=8, leading=10)
    story = []
    for kind, text in doc["lines"]:
        story.append(Paragraph(text, styles[kind]) if text else Spacer(1, 6 * mm))
    story += [Spacer(1, 10 * mm), Paragraph(FOOTER, small)]
    pdf = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=25 * mm, rightMargin=15 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title=doc["about"], author="Сибирский Вайб — тестовые данные", invariant=1,
    )
    pdf.build(story)


def write_expected() -> None:
    expected = {
        "comment": (
            "Ожидаемый результат для каждого тестового документа. fields — что должно получиться "
            "после идеального извлечения полей; issues — правила из rules.json, которые обязаны "
            "сработать, в порядке правил. referenceDate — «сегодня» для DATE_NOT_FUTURE."
        ),
        "referenceDate": REFERENCE_DATE.isoformat(),
        "documents": [
            {
                "file": f"documents/{d['file']}",
                "type": d["type"],
                "about": d["about"],
                "fields": d["fields"],
                "issues": d["issues"],
            }
            for d in DOCS
        ],
    }
    (TESTDATA / "expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )


# ---------------------------------------------------------------------------
# docs/RULES.md
# ---------------------------------------------------------------------------

KIND = {"LEGAL": "Закон", "INTERNAL_POLICY": "Регламент компании", "PRODUCT_RULE": "Правило продукта"}
SEVERITY = {"BLOCKER": "Блокирует отправку", "WARNING": "Предупреждение", "INFO": "Подсказка"}
CHECK = {
    "REQUIRED": "Заполнено",
    "DATE_FORMAT": "Формат даты",
    "DATE_NOT_FUTURE": "Дата не в будущем",
    "MATCHES_PATTERN": "Соответствует шаблону",
    "MIN_LENGTH": "Не короче",
    "ONE_OF": "Одно из значений",
}


def cell(text) -> str:
    """Экранирует то, что ломает таблицу Markdown: «|» делит столбцы, «<» GitHub принимает за тег."""
    if text in (None, ""):
        return "—"
    return str(text).replace("|", "\\|").replace("<", "&lt;").replace("\n", " ")


def write_rules_md() -> None:
    sources = RULES["sources"]
    out = [
        "# Правила проверки документов",
        "",
        "Сгенерировано из [`testdata/rules.json`](../testdata/rules.json) скриптом",
        "`testdata/generate.py` — правьте JSON и перезапускайте скрипт, а не этот файл.",
        f"Источники сверены {datetime.fromisoformat(RULES['sourceCheckedAt']).strftime('%d.%m.%Y')}.",
        "",
        "Модель только извлекает поля, правила проверяет обычный код — `RuleEngine`",
        "(CLAUDE.md, «Модель извлекает, код решает»). Пустое поле проверяет только",
        "«Заполнено»: остальные проверки его пропускают, иначе одно отсутствующее поле",
        "давало бы несколько одинаковых замечаний.",
        "",
        "**Виды правил** — то, что жюри увидит у каждого замечания:",
        "- **Закон** — норма федерального закона, со ссылкой на статью и редакцию.",
        "  Только у заявки на меру поддержки: она уходит наружу, цена ошибки выше.",
        "- **Регламент компании** — внутренний стандарт. У служебной записки это ГОСТ Р 7.0.97-2016,",
        "  принятый компанией: для коммерческой организации ГОСТ добровольный, поэтому это регламент,",
        "  а не закон. У заявления на отпуск — положение демо-компании; оно вымышленное,",
        "  как и сама демо-компания: в настоящей компании здесь её локальный акт.",
        "- **Правило продукта** — наше допущение, без которого проверка теряет смысл.",
        "",
    ]
    for doc_type in RULES["types"]:
        out += [f"## {doc_type['name']} (`{doc_type['code']}`)", "", "### Поля", "",
                "| Поле | Тип | Название | Подсказка для извлечения |", "|---|---|---|---|"]
        for f in doc_type["fields"]:
            out.append(f"| `{f['name']}` | {f['type']} | {cell(f['label'])} | {cell(f['hint'])} |")
        out += ["", "### Правила", "",
                "| Код | Поле | Проверка | Параметр | Вид | Важность | Источник | Текст замечания |",
                "|---|---|---|---|---|---|---|---|"]
        for r in doc_type["rules"]:
            src = sources[r["source"]]
            short = {
                "GOST": "ГОСТ Р 7.0.97-2016",
                "VACATION_POLICY": "Положение об отпусках (демо)",
                "PRODUCT": "Правило продукта",
                "LAW_209_ART_4_1": "209-ФЗ",
                "LAW_209_ART_14": "209-ФЗ",
            }[r["source"]]
            ref = f"{short}, {r['sourceRef']}" if r["sourceRef"] else short
            link = f"[{cell(ref)}]({src['url']})" if src["url"] else cell(ref)
            # Внутри кода в таблице GitHub понимает только экранированный «\|».
            param = "`" + r["expected"].replace("|", "\\|") + "`" if r["expected"] else "—"
            out.append(
                f"| {r['code']} | `{r['field']}` | {CHECK[r['check']]} | {param} | {KIND[r['kind']]} "
                f"| {SEVERITY[r['severity']]} | {link} | {cell(r['description'])} |"
            )
        out.append("")
    out += ["## Полные названия источников", ""]
    for key, src in sources.items():
        url = f" — {src['url']}" if src["url"] else ""
        out.append(f"- **{key}**: {src['title']}{url}")
    out += [
        "",
        "## Тестовые документы",
        "",
        "Лежат в [`testdata/documents`](../testdata/documents), ожидаемые поля и замечания —",
        "в [`testdata/expected.json`](../testdata/expected.json).",
        "",
        "| Файл | Тип | Что в нём | Какие правила срабатывают |",
        "|---|---|---|---|",
    ]
    for d in DOCS:
        issues = ", ".join(d["issues"]) or "никакие"
        out.append(f"| `{d['file']}` | `{d['type']}` | {cell(d['about'])} | {issues} |")
    out.append("")
    (ROOT / "docs" / "RULES.md").write_text("\n".join(out), encoding="utf-8", newline="\n")


def main() -> None:
    self_check()
    DOCUMENTS.mkdir(parents=True, exist_ok=True)
    font = register_font()
    for doc in DOCS:
        path = DOCUMENTS / doc["file"]
        if path.suffix == ".docx":
            write_docx(doc, path)
        else:
            write_pdf(doc, path, font)
    write_expected()
    write_rules_md()
    print(f"Готово: {len(DOCS)} документов, expected.json, docs/RULES.md")


if __name__ == "__main__":
    main()
